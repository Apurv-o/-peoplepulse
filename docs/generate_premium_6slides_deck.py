import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

def build_presentation(output_path="PeoplePulse_Hackathon_Premium_6Slides.pptx"):
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # Premium Modern SaaS Color Palette (Warm Light / Minimalist)
    C_CANVAS_BG   = RGBColor(248, 250, 252)  # #F8FAFC Warm Light Canvas
    C_CARD_BG     = RGBColor(255, 255, 255)  # #FFFFFF Pure White
    C_CARD_BORDER = RGBColor(226, 232, 240)  # #E2E8F0 Subtle Border
    C_TITLE       = RGBColor(15, 23, 42)     # #0F172A Deep Slate Title
    C_BODY        = RGBColor(51, 65, 85)     # #334155 Slate Body
    C_MUTED       = RGBColor(100, 116, 139)  # #64748B Subtitle / Meta
    C_PRIMARY     = RGBColor(37, 99, 235)    # #2563EB PeoplePulse Blue
    C_PRIMARY_TINT= RGBColor(239, 246, 255)  # #EFF6FF Light Blue BG
    C_AI_PURPLE   = RGBColor(124, 58, 237)   # #7C3AED AI Decision Purple
    C_PURPLE_TINT = RGBColor(245, 243, 255)  # #F5F3FF Light Purple BG
    C_GREEN       = RGBColor(16, 185, 129)   # #10B981 Action/Success Green
    C_GREEN_TINT  = RGBColor(236, 253, 245)  # #ECFDF5 Light Green BG
    C_AMBER       = RGBColor(245, 158, 11)   # #F59E0B Warning Amber
    C_AMBER_TINT  = RGBColor(254, 243, 199)  # #FEF3C7 Light Amber BG
    C_RED_TINT    = RGBColor(254, 242, 242)  # #FEF2F2 Soft Red

    def apply_slide_canvas(slide):
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
        bg.fill.solid()
        bg.fill.fore_color.rgb = C_CANVAS_BG
        bg.line.fill.background()
        return bg

    def add_header(slide, category, title, slide_num, subtitle=""):
        # Category / Event Tag Pill
        cat_box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(0.4), Inches(4.5), Inches(0.32))
        cat_box.fill.solid()
        cat_box.fill.fore_color.rgb = C_PRIMARY_TINT
        cat_box.line.color.rgb = RGBColor(191, 219, 254)
        cat_box.line.width = Pt(1)
        tf_cat = cat_box.text_frame
        tf_cat.word_wrap = True
        tf_cat.vertical_anchor = MSO_ANCHOR.MIDDLE
        p_c = tf_cat.paragraphs[0]
        p_c.text = f"TECH ZEPHYR 4.0 | IIT BHUBANESWAR • {category.upper()}"
        p_c.font.size = Pt(8.5)
        p_c.font.bold = True
        p_c.font.color.rgb = C_PRIMARY
        p_c.alignment = PP_ALIGN.CENTER

        # Slide Number Pill
        num_box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(11.4), Inches(0.4), Inches(1.15), Inches(0.32))
        num_box.fill.solid()
        num_box.fill.fore_color.rgb = RGBColor(241, 245, 249)
        num_box.line.color.rgb = C_CARD_BORDER
        num_box.line.width = Pt(0.75)
        tf_n = num_box.text_frame
        tf_n.vertical_anchor = MSO_ANCHOR.MIDDLE
        p_n = tf_n.paragraphs[0]
        p_n.text = f"Slide {slide_num} of 6"
        p_n.font.size = Pt(8.5)
        p_n.font.bold = True
        p_n.font.color.rgb = C_MUTED
        p_n.alignment = PP_ALIGN.CENTER

        # Title
        t_box = slide.shapes.add_textbox(Inches(0.76), Inches(0.78), Inches(11.8), Inches(0.65))
        tf_t = t_box.text_frame
        tf_t.word_wrap = True
        p_t = tf_t.paragraphs[0]
        p_t.text = title
        p_t.font.size = Pt(22)
        p_t.font.bold = True
        p_t.font.color.rgb = C_TITLE

        if subtitle:
            p_sub = tf_t.add_paragraph()
            p_sub.text = subtitle
            p_sub.font.size = Pt(11)
            p_sub.font.color.rgb = C_MUTED
            p_sub.space_before = Pt(2)

    def draw_card(slide, left, top, width, height, bg_color=C_CARD_BG, border_color=C_CARD_BORDER, border_width=1):
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        card.fill.solid()
        card.fill.fore_color.rgb = bg_color
        card.line.color.rgb = border_color
        card.line.width = Pt(border_width)
        return card

    # =========================================================================
    # SLIDE 1: PROBLEM & VISION
    # =========================================================================
    s1 = prs.slides.add_slide(blank_layout)
    apply_slide_canvas(s1)
    add_header(s1, "Problem & Vision", "PeoplePulse", "1", "From Employee Feedback to Intelligent HR Action")

    # Flow Diagram: Employee Check-ins -> Dashboards -> Manual Analysis -> Delayed Action
    draw_card(s1, Inches(0.8), Inches(1.5), Inches(11.73), Inches(0.85), bg_color=C_CARD_BG)
    
    flow_steps = [
        ("1. Daily Check-ins", "Short pulses submitted", C_PRIMARY),
        ("2. Static Dashboards", "Raw charts & tables", C_MUTED),
        ("3. Manual Analysis", "HR spreadsheets & delays", C_AMBER),
        ("4. Delayed Action", "Burnout escalates silently", RGBColor(239, 68, 68))
    ]
    step_w = Inches(2.6)
    for i, (st, desc, col) in enumerate(flow_steps):
        s_left = Inches(0.95) + i * Inches(2.9)
        step_box = s1.shapes.add_textbox(s_left, Inches(1.55), step_w, Inches(0.75))
        tf = step_box.text_frame
        tf.word_wrap = True
        p1 = tf.paragraphs[0]
        p1.text = st
        p1.font.size = Pt(11)
        p1.font.bold = True
        p1.font.color.rgb = col
        p2 = tf.add_paragraph()
        p2.text = desc
        p2.font.size = Pt(9)
        p2.font.color.rgb = C_MUTED
        
        # Arrow
        if i < 3:
            arr = s1.shapes.add_textbox(s_left + step_w - Inches(0.1), Inches(1.68), Inches(0.4), Inches(0.4))
            arr_tf = arr.text_frame
            p_arr = arr_tf.paragraphs[0]
            p_arr.text = "➔"
            p_arr.font.size = Pt(12)
            p_arr.font.color.rgb = C_MUTED

    # Left Column: 3 Pain Points Cards
    left_w = Inches(5.65)
    card_h = Inches(1.15)
    pains = [
        ("Slow Action Cycle", "Employee feedback is collected regularly, but action response is slow and bureaucratic.", C_AMBER),
        ("Invisible Burnout", "Managers struggle to identify burnout signals before critical team members resign.", RGBColor(239, 68, 68)),
        ("Passive Dashboards", "Traditional HR tools only display passive insights—they cannot execute interventions.", C_AI_PURPLE)
    ]
    for idx, (p_title, p_desc, col) in enumerate(pains):
        top_pos = Inches(2.55) + idx * Inches(1.3)
        draw_card(s1, Inches(0.8), top_pos, left_w, card_h)
        # Indicator bar on left
        bar = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), top_pos, Inches(0.12), card_h)
        bar.fill.solid()
        bar.fill.fore_color.rgb = col
        bar.line.fill.background()

        tb = s1.shapes.add_textbox(Inches(1.05), top_pos + Inches(0.15), left_w - Inches(0.35), card_h - Inches(0.3))
        tf = tb.text_frame
        tf.word_wrap = True
        p0 = tf.paragraphs[0]
        p0.text = p_title
        p0.font.size = Pt(12)
        p0.font.bold = True
        p0.font.color.rgb = C_TITLE
        p1 = tf.add_paragraph()
        p1.text = p_desc
        p1.font.size = Pt(10)
        p1.font.color.rgb = C_BODY
        p1.space_before = Pt(3)

    # Right Column: Dashboard Mockup Visual
    right_left = Inches(6.7)
    right_w = Inches(5.83)
    right_h = Inches(3.75)
    draw_card(s1, right_left, Inches(2.55), right_w, right_h, bg_color=C_CARD_BG)
    
    # Header inside mockup
    mock_hdr = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, right_left + Inches(0.2), Inches(2.72), right_w - Inches(0.4), Inches(0.45))
    mock_hdr.fill.solid()
    mock_hdr.fill.fore_color.rgb = RGBColor(241, 245, 249)
    mock_hdr.line.color.rgb = C_CARD_BORDER
    tf_mh = mock_hdr.text_frame
    tf_mh.vertical_anchor = MSO_ANCHOR.MIDDLE
    p_mh = tf_mh.paragraphs[0]
    p_mh.text = "⚡ PeoplePulse Live Workplace Radar — Active Tenant Analytics"
    p_mh.font.size = Pt(9.5)
    p_mh.font.bold = True
    p_mh.font.color.rgb = C_PRIMARY

    # Mock metrics grid inside
    metrics = [
        ("84%", "Daily Participation", C_PRIMARY),
        ("38%", "Burnout Risk (CS)", RGBColor(239, 68, 68)),
        ("7.8/10", "Avg Engagement", C_GREEN)
    ]
    for mi, (m_val, m_lbl, m_c) in enumerate(metrics):
        m_box = draw_card(s1, right_left + Inches(0.2) + mi * Inches(1.8), Inches(3.3), Inches(1.65), Inches(1.15), bg_color=C_CANVAS_BG)
        mtb = s1.shapes.add_textbox(right_left + Inches(0.25) + mi * Inches(1.8), Inches(3.38), Inches(1.55), Inches(1.0))
        mtf = mtb.text_frame
        p_v = mtf.paragraphs[0]
        p_v.text = m_val
        p_v.font.size = Pt(18)
        p_v.font.bold = True
        p_v.font.color.rgb = m_c
        p_l = mtf.add_paragraph()
        p_l.text = m_lbl
        p_l.font.size = Pt(8.5)
        p_l.font.bold = True
        p_l.font.color.rgb = C_MUTED

    # Team breakdown mock rows
    teams_data = [
        ("Customer Success", "68% Engagement", "High Workload Alert", RGBColor(239, 68, 68), C_RED_TINT),
        ("Product & Eng", "88% Engagement", "Healthy Trend", C_GREEN, C_GREEN_TINT),
        ("Marketing & Growth", "79% Engagement", "Stable", C_PRIMARY, C_PRIMARY_TINT)
    ]
    for ti, (tname, tscore, tbadge, tc, tbg) in enumerate(teams_data):
        t_row = draw_card(s1, right_left + Inches(0.2), Inches(4.6) + ti * Inches(0.52), right_w - Inches(0.4), Inches(0.42), bg_color=tbg, border_color=C_CARD_BORDER)
        t_tb = s1.shapes.add_textbox(right_left + Inches(0.3), Inches(4.62) + ti * Inches(0.52), right_w - Inches(0.6), Inches(0.38))
        ttf = t_tb.text_frame
        tp = ttf.paragraphs[0]
        tp.text = f"{tname}  |  {tscore}  |  ● {tbadge}"
        tp.font.size = Pt(9)
        tp.font.bold = True
        tp.font.color.rgb = tc

    # Bottom Closing Statement Banner
    bot_card = draw_card(s1, Inches(0.8), Inches(6.45), Inches(11.73), Inches(0.65), bg_color=C_PRIMARY_TINT, border_color=RGBColor(191, 219, 254), border_width=1.5)
    bot_tb = s1.shapes.add_textbox(Inches(0.9), Inches(6.48), Inches(11.53), Inches(0.55))
    bot_tf = bot_tb.text_frame
    bot_tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    bot_p = bot_tf.paragraphs[0]
    bot_p.text = "PeoplePulse transforms daily employee feedback into meaningful organizational action using Agentic AI."
    bot_p.font.size = Pt(11)
    bot_p.font.bold = True
    bot_p.font.color.rgb = C_PRIMARY
    bot_p.alignment = PP_ALIGN.CENTER

    # =========================================================================
    # SLIDE 2: MEET PULSEAGENT
    # =========================================================================
    s2 = prs.slides.add_slide(blank_layout)
    apply_slide_canvas(s2)
    add_header(s2, "Autonomous HR Copilot", "Meet PulseAgent — Your Autonomous HR Copilot", "2", "PulseAgent is not just a chatbot—it is an autonomous agent executing the full ReAct loop.")

    # Intro sub-banner
    intro_card = draw_card(s2, Inches(0.8), Inches(1.5), Inches(11.73), Inches(0.7), bg_color=C_CARD_BG)
    itb = s2.shapes.add_textbox(Inches(0.95), Inches(1.55), Inches(11.4), Inches(0.6))
    itf = itb.text_frame
    itf.word_wrap = True
    ip = itf.paragraphs[0]
    ip.text = "PulseAgent is not just a chatbot. It observes organization health, decides the next best action, uses real tools, evaluates results, and adapts when something fails."
    ip.font.size = Pt(11)
    ip.font.color.rgb = C_BODY

    # 4 Visual Cards: Observe, Decide, Act, Adapt
    cards_data = [
        ("Observe 👁", "Continuous Health Monitor", "Understands engagement, stress, workload, and sentiment through real employee check-in telemetry.", C_PRIMARY, C_PRIMARY_TINT),
        ("Decide 🧠", "Autonomous Reasoning", "Chooses the appropriate intervention using authorized policy frameworks and Gemini intelligence.", C_AI_PURPLE, C_PURPLE_TINT),
        ("Act 🔧", "Deterministic Tool Execution", "Executes approved HR tools—diagnosing teams, dispatching surveys, and sending briefs.", C_GREEN, C_GREEN_TINT),
        ("Adapt 🔄", "Self-Correcting Strategy", "Evaluates live outcomes, detects channel failures, and switches strategies automatically.", C_AMBER, C_AMBER_TINT)
    ]
    card_w = Inches(2.78)
    card_h = Inches(3.8)
    for ci, (c_title, c_sub, c_body, c_acc, c_bg) in enumerate(cards_data):
        c_left = Inches(0.8) + ci * Inches(2.98)
        draw_card(s2, c_left, Inches(2.35), card_w, card_h, bg_color=C_CARD_BG)
        
        # Color top banner
        top_strip = s2.shapes.add_shape(MSO_SHAPE.RECTANGLE, c_left, Inches(2.35), card_w, Inches(0.12))
        top_strip.fill.solid()
        top_strip.fill.fore_color.rgb = c_acc
        top_strip.line.fill.background()

        # Icon / Title badge
        badge = draw_card(s2, c_left + Inches(0.2), Inches(2.65), card_w - Inches(0.4), Inches(0.55), bg_color=c_bg, border_color=C_CARD_BORDER)
        btb = s2.shapes.add_textbox(c_left + Inches(0.25), Inches(2.7), card_w - Inches(0.5), Inches(0.45))
        btf = btb.text_frame
        bp = btf.paragraphs[0]
        bp.text = c_title
        bp.font.size = Pt(13)
        bp.font.bold = True
        bp.font.color.rgb = c_acc

        # Subtitle
        stb = s2.shapes.add_textbox(c_left + Inches(0.2), Inches(3.3), card_w - Inches(0.4), Inches(0.4))
        stf = stb.text_frame
        sp = stf.paragraphs[0]
        sp.text = c_sub
        sp.font.size = Pt(9.5)
        sp.font.bold = True
        sp.font.color.rgb = C_TITLE

        # Body text
        dtb = s2.shapes.add_textbox(c_left + Inches(0.2), Inches(3.75), card_w - Inches(0.4), Inches(1.4))
        dtf = dtb.text_frame
        dtf.word_wrap = True
        dp = dtf.paragraphs[0]
        dp.text = c_body
        dp.font.size = Pt(10)
        dp.font.color.rgb = C_BODY

        # Feature bullet pill at bottom of card
        ftb = draw_card(s2, c_left + Inches(0.2), Inches(5.35), card_w - Inches(0.4), Inches(0.6), bg_color=C_CANVAS_BG)
        ft_tb = s2.shapes.add_textbox(c_left + Inches(0.25), Inches(5.38), card_w - Inches(0.5), Inches(0.5))
        ftf = ft_tb.text_frame
        ftf.word_wrap = True
        ftp = ftf.paragraphs[0]
        ftp.text = f"✓ Tenant Protected\n✓ Authorized Action"
        ftp.font.size = Pt(8.5)
        ftp.font.bold = True
        ftp.font.color.rgb = C_MUTED

    # Bottom statement
    s2_bot = draw_card(s2, Inches(0.8), Inches(6.35), Inches(11.73), Inches(0.7), bg_color=C_PRIMARY_TINT, border_color=RGBColor(191, 219, 254), border_width=1.5)
    s2_tb = s2.shapes.add_textbox(Inches(0.9), Inches(6.4), Inches(11.53), Inches(0.6))
    s2_tf = s2_tb.text_frame
    s2_tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    s2_p = s2_tf.paragraphs[0]
    s2_p.text = "An AI teammate that helps HR move from insight to action."
    s2_p.font.size = Pt(12)
    s2_p.font.bold = True
    s2_p.font.color.rgb = C_PRIMARY
    s2_p.alignment = PP_ALIGN.CENTER

    # =========================================================================
    # SLIDE 3: SYSTEM ARCHITECTURE
    # =========================================================================
    s3 = prs.slides.add_slide(blank_layout)
    apply_slide_canvas(s3)
    add_header(s3, "System Architecture", "How PulseAgent Works", "3", "Enterprise-grade architecture built for speed, tenant security, and observable execution.")

    # Top Flow Pipeline: Simple 30-second judge comprehension flow
    pipe_card = draw_card(s3, Inches(0.8), Inches(1.5), Inches(11.73), Inches(0.7), bg_color=C_CARD_BG)
    pipe_tb = s3.shapes.add_textbox(Inches(0.9), Inches(1.53), Inches(11.53), Inches(0.6))
    ptf = pipe_tb.text_frame
    ptf.vertical_anchor = MSO_ANCHOR.MIDDLE
    pp = ptf.paragraphs[0]
    pp.text = "Employee / Manager  ➔  PeoplePulse Web App  ➔  PulseAgent Orchestrator  ➔  Tool Registry  ➔  Supabase & AI  ➔  Observe  ➔  Evaluate  ➔  Adapt  ➔  HR Outcome"
    pp.font.size = Pt(8.5)
    pp.font.bold = True
    pp.font.color.rgb = C_PRIMARY
    pp.alignment = PP_ALIGN.CENTER

    # 4 Architecture Columns: Frontend, Agent Layer, Data Layer, Security
    arch_cols = [
        ("Frontend", "Client Experience", [
            "React 18 & Vite SPA",
            "Employee Pulse Dashboard",
            "Manager Health Console",
            "Admin Governance Portal",
            "PulseAgent Interactive Chat"
        ], C_PRIMARY, C_PRIMARY_TINT),
        ("Agent Layer", "Autonomous ReAct Loop", [
            "Goal Planner & Parser",
            "Gemini Decision Engine",
            "Deterministic Tool Executor",
            "Evaluation Engine",
            "Dynamic Adaptation Logic",
            "Immutable Audit Trail"
        ], C_AI_PURPLE, C_PURPLE_TINT),
        ("Data Layer", "Supabase PostgreSQL", [
            "Daily Check-in Telemetry",
            "Team Hierarchies & Org Map",
            "Engagement & Stress Metrics",
            "Sentiment Analysis Store",
            "Survey Question Bank"
        ], C_GREEN, C_GREEN_TINT),
        ("Security & Trust", "Enterprise Governance", [
            "Strict Role-Based Access (RBAC)",
            "Tenant Isolation (RLS)",
            "Anonymous Privacy Shield",
            "Server-Side Secret Guards",
            "Authorized Tool Allowlist"
        ], C_AMBER, C_AMBER_TINT)
    ]

    col_w = Inches(2.78)
    col_h = Inches(4.35)
    for ai, (col_name, col_sub, col_items, col_acc, col_bg) in enumerate(arch_cols):
        col_left = Inches(0.8) + ai * Inches(2.98)
        draw_card(s3, col_left, Inches(2.35), col_w, col_h, bg_color=C_CARD_BG)

        # Header bar
        top_bar = s3.shapes.add_shape(MSO_SHAPE.RECTANGLE, col_left, Inches(2.35), col_w, Inches(0.12))
        top_bar.fill.solid()
        top_bar.fill.fore_color.rgb = col_acc
        top_bar.line.fill.background()

        # Title
        hdr_b = draw_card(s3, col_left + Inches(0.18), Inches(2.6), col_w - Inches(0.36), Inches(0.62), bg_color=col_bg, border_color=C_CARD_BORDER)
        ht = s3.shapes.add_textbox(col_left + Inches(0.22), Inches(2.63), col_w - Inches(0.44), Inches(0.55))
        htf = ht.text_frame
        hp1 = htf.paragraphs[0]
        hp1.text = col_name
        hp1.font.size = Pt(12)
        hp1.font.bold = True
        hp1.font.color.rgb = col_acc
        hp2 = htf.add_paragraph()
        hp2.text = col_sub
        hp2.font.size = Pt(8.5)
        hp2.font.color.rgb = C_MUTED

        # Items
        it_tb = s3.shapes.add_textbox(col_left + Inches(0.18), Inches(3.35), col_w - Inches(0.36), Inches(3.2))
        it_tf = it_tb.text_frame
        it_tf.word_wrap = True
        for ii, item in enumerate(col_items):
            p = it_tf.paragraphs[0] if ii == 0 else it_tf.add_paragraph()
            p.text = f"• {item}"
            p.font.size = Pt(9.5)
            p.font.color.rgb = C_BODY
            p.space_before = Pt(4)

    # Bottom caption
    s3_bot = s3.shapes.add_textbox(Inches(0.8), Inches(6.8), Inches(11.73), Inches(0.4))
    s3_tf = s3_bot.text_frame
    s3_p = s3_tf.paragraphs[0]
    s3_p.text = "Engineered for clarity: Judges can understand the entire end-to-end stack in under 30 seconds."
    s3_p.font.size = Pt(9.5)
    s3_p.font.bold = True
    s3_p.font.color.rgb = C_MUTED
    s3_p.alignment = PP_ALIGN.CENTER

    # =========================================================================
    # SLIDE 4: AGENTIC WORKFLOW (HERO SLIDE)
    # =========================================================================
    s4 = prs.slides.add_slide(blank_layout)
    apply_slide_canvas(s4)
    add_header(s4, "Hero Slide • Autonomous ReAct Loop", "Goal ➔ Decision ➔ Action ➔ Evaluation ➔ Adaptation ➔ Outcome", "4", "Live execution walkthrough: Customer Success burnout investigation with real tool dispatching.")

    # 6 Step Cards arranged horizontally across 11.73 inches
    steps_data = [
        ("GOAL", "🎯", "Investigate Burnout", "“Investigate Customer Success team burnout.”", C_PRIMARY, C_PRIMARY_TINT),
        ("DECISION", "🧠", "Inspect Health", "PulseAgent decides to inspect organization and team health metrics.", C_AI_PURPLE, C_PURPLE_TINT),
        ("ACTION", "🔧", "Call Approved Tools", "Calls approved tools:\n• get_organization_metrics\n• diagnose_team_health", C_GREEN, C_GREEN_TINT),
        ("OBSERVATION", "👁", "Live Telemetry", "Receives live engagement (68%), stress (78%), and participation data.", C_PRIMARY, C_PRIMARY_TINT),
        ("EVALUATION", "⚖️", "Intervention Test", "Determines that proactive intervention is critically required.", C_AMBER, C_AMBER_TINT),
        ("ADAPTATION", "🔄", "Execute Response", "Creates a targeted adaptive survey and generates manager brief.", C_AI_PURPLE, C_PURPLE_TINT),
    ]

    s_w = Inches(1.82)
    s_h = Inches(3.6)
    for si, (s_label, s_icon, s_heading, s_desc, s_color, s_bg) in enumerate(steps_data):
        s_left = Inches(0.8) + si * Inches(1.98)
        draw_card(s4, s_left, Inches(1.5), s_w, s_h, bg_color=C_CARD_BG)

        # Top tag
        tag = s4.shapes.add_shape(MSO_SHAPE.RECTANGLE, s_left, Inches(1.5), s_w, Inches(0.1))
        tag.fill.solid()
        tag.fill.fore_color.rgb = s_color
        tag.line.fill.background()

        # Step Pill
        pill = draw_card(s4, s_left + Inches(0.12), Inches(1.7), s_w - Inches(0.24), Inches(0.42), bg_color=s_bg, border_color=C_CARD_BORDER)
        ptb = s4.shapes.add_textbox(s_left + Inches(0.14), Inches(1.72), s_w - Inches(0.28), Inches(0.38))
        ptf = ptb.text_frame
        pp = ptf.paragraphs[0]
        pp.text = f"{s_icon} {s_label}"
        pp.font.size = Pt(8.5)
        pp.font.bold = True
        pp.font.color.rgb = s_color

        # Heading
        htb = s4.shapes.add_textbox(s_left + Inches(0.12), Inches(2.2), s_w - Inches(0.24), Inches(0.45))
        htf = htb.text_frame
        hp = htf.paragraphs[0]
        hp.text = s_heading
        hp.font.size = Pt(9.5)
        hp.font.bold = True
        hp.font.color.rgb = C_TITLE

        # Desc
        dtb = s4.shapes.add_textbox(s_left + Inches(0.12), Inches(2.7), s_w - Inches(0.24), Inches(2.2))
        dtf = dtb.text_frame
        dtf.word_wrap = True
        dp = dtf.paragraphs[0]
        dp.text = s_desc
        dp.font.size = Pt(8.5)
        dp.font.color.rgb = C_BODY

        # Connector arrow
        if si < 5:
            carr = s4.shapes.add_textbox(s_left + s_w - Inches(0.08), Inches(3.1), Inches(0.28), Inches(0.3))
            catf = carr.text_frame
            cap = catf.paragraphs[0]
            cap.text = "➔"
            cap.font.size = Pt(11)
            cap.font.color.rgb = C_PRIMARY

    # Final Outcome Card Banner Below
    out_card = draw_card(s4, Inches(0.8), Inches(5.25), Inches(11.73), Inches(0.95), bg_color=C_GREEN_TINT, border_color=RGBColor(167, 243, 208), border_width=1.5)
    otb = s4.shapes.add_textbox(Inches(1.0), Inches(5.3), Inches(11.33), Inches(0.85))
    otf = otb.text_frame
    otf.word_wrap = True
    op1 = otf.paragraphs[0]
    op1.text = "FINAL OUTCOME: Measurable HR Resolution"
    op1.font.size = Pt(11)
    op1.font.bold = True
    op1.font.color.rgb = C_GREEN
    op2 = otf.add_paragraph()
    op2.text = "Managers receive actionable 1:1 talking points while employee privacy remains 100% protected through aggregation thresholds. Follow-up pulse survey deployed to track recovery."
    op2.font.size = Pt(9.5)
    op2.font.color.rgb = C_BODY
    op2.space_before = Pt(2)

    # Bottom caption
    s4_bot = s4.shapes.add_textbox(Inches(0.8), Inches(6.35), Inches(11.73), Inches(0.45))
    s4_tf = s4_bot.text_frame
    s4_p = s4_tf.paragraphs[0]
    s4_p.text = "Every step is authorized, observable and auditable. (Zero exposed chain-of-thought; strictly safe reasoning summaries)."
    s4_p.font.size = Pt(9.5)
    s4_p.font.bold = True
    s4_p.font.color.rgb = C_MUTED
    s4_p.alignment = PP_ALIGN.CENTER

    # =========================================================================
    # SLIDE 5: LIVE DEMONSTRATION
    # =========================================================================
    s5 = prs.slides.add_slide(blank_layout)
    apply_slide_canvas(s5)
    add_header(s5, "Live Demonstration", "Live Demo Scenario & Execution Trace", "5", "Watch PulseAgent handle an urgent prompt live in production with real tool execution.")

    # Top Prompt Banner
    prompt_card = draw_card(s5, Inches(0.8), Inches(1.5), Inches(11.73), Inches(0.72), bg_color=C_PRIMARY_TINT, border_color=RGBColor(191, 219, 254), border_width=1.5)
    pr_tb = s5.shapes.add_textbox(Inches(1.0), Inches(1.55), Inches(11.33), Inches(0.6))
    pr_tf = pr_tb.text_frame
    pr_p = pr_tf.paragraphs[0]
    pr_p.text = "USER PROMPT: “Investigate the Customer Success team's health and take the appropriate action.”"
    pr_p.font.size = Pt(11.5)
    pr_p.font.bold = True
    pr_p.font.color.rgb = C_PRIMARY

    # Left: Execution Timeline Stages
    left_x = Inches(0.8)
    left_width = Inches(5.2)
    draw_card(s5, left_x, Inches(2.35), left_width, Inches(4.3), bg_color=C_CARD_BG)

    ltb = s5.shapes.add_textbox(left_x + Inches(0.2), Inches(2.45), left_width - Inches(0.4), Inches(4.1))
    ltf = ltb.text_frame
    ltf.word_wrap = True

    timeline_items = [
        ("GOAL", "Investigate team wellbeing", C_PRIMARY),
        ("DECISION", "Analyze engagement and stress telemetry", C_AI_PURPLE),
        ("ACTION", "Run diagnostic tools on Customer Success", C_GREEN),
        ("OBSERVATION", "Identify low engagement (68%) & capacity strain", C_PRIMARY),
        ("EVALUATION", "Choose best proactive intervention path", C_AMBER),
        ("ADAPTATION", "Deploy adaptive survey + generate manager brief", C_AI_PURPLE),
        ("FINAL RESULT", "Clear HR guidance + persistent Supabase actions", C_GREEN)
    ]
    for ti, (tag, detail, col) in enumerate(timeline_items):
        tp1 = ltf.paragraphs[0] if ti == 0 else ltf.add_paragraph()
        tp1.text = f"{tag} ➔ {detail}"
        tp1.font.size = Pt(9.5)
        tp1.font.bold = True
        tp1.font.color.rgb = col
        tp1.space_before = Pt(4) if ti > 0 else Pt(0)

    # Right: Mock Execution Trace (Inspired by PulseAgent UI)
    right_x = Inches(6.2)
    right_width = Inches(6.33)
    draw_card(s5, right_x, Inches(2.35), right_width, Inches(4.3), bg_color=C_CARD_BG)

    # Header for Mock Trace
    t_hdr = draw_card(s5, right_x + Inches(0.18), Inches(2.48), right_width - Inches(0.36), Inches(0.38), bg_color=RGBColor(241, 245, 249))
    th_tb = s5.shapes.add_textbox(right_x + Inches(0.25), Inches(2.5), right_width - Inches(0.5), Inches(0.32))
    th_tf = th_tb.text_frame
    th_p = th_tf.paragraphs[0]
    th_p.text = "TERMINAL • PulseAgent Autonomous Trace Stream"
    th_p.font.size = Pt(8.5)
    th_p.font.bold = True
    th_p.font.color.rgb = C_PRIMARY

    # Trace log entries
    trace_entries = [
        ("[TOOL EXEC]", "get_organization_metrics(organization_id)", "Checked 142 check-ins across 5 teams (Org avg: 74%).", C_PRIMARY_TINT, C_PRIMARY),
        ("[TOOL EXEC]", "diagnose_team_health(team='Customer Success')", "CS Engagement: 68% | Stress: High | Workload: Critical.", C_AMBER_TINT, C_AMBER),
        ("[EVALUATION]", "Intervention Threshold Reached", "Score falls below 70% threshold. Immediate action triggered.", C_PURPLE_TINT, C_AI_PURPLE),
        ("[ACTION]", "dispatch_adaptive_survey(target='Customer Success')", "Inserted: “Do you have bandwidth for sprint priorities?”", C_GREEN_TINT, C_GREEN),
        ("[ACTION]", "trigger_manager_action_brief(team='Customer Success')", "Coaching talking points dispatched to team lead.", C_GREEN_TINT, C_GREEN),
    ]

    for ei, (badge, title, desc, bgc, textc) in enumerate(trace_entries):
        ey = Inches(2.98) + ei * Inches(0.68)
        draw_card(s5, right_x + Inches(0.18), ey, right_width - Inches(0.36), Inches(0.58), bg_color=bgc, border_color=C_CARD_BORDER)
        etb = s5.shapes.add_textbox(right_x + Inches(0.24), ey + Inches(0.04), right_width - Inches(0.48), Inches(0.5))
        etf = etb.text_frame
        etf.word_wrap = True
        ep1 = etf.paragraphs[0]
        ep1.text = f"{badge} {title}"
        ep1.font.size = Pt(8.5)
        ep1.font.bold = True
        ep1.font.color.rgb = textc
        ep2 = etf.add_paragraph()
        ep2.text = desc
        ep2.font.size = Pt(8)
        ep2.font.color.rgb = C_BODY

    # Small Bottom Note
    s5_bot = s5.shapes.add_textbox(Inches(0.8), Inches(6.8), Inches(11.73), Inches(0.4))
    s5_tf = s5_bot.text_frame
    s5_p = s5_tf.paragraphs[0]
    s5_p.text = "Live demonstration uses real application data and actual tool execution. No simulated recordings."
    s5_p.font.size = Pt(9.5)
    s5_p.font.bold = True
    s5_p.font.color.rgb = C_PRIMARY
    s5_p.alignment = PP_ALIGN.CENTER

    # =========================================================================
    # SLIDE 6: IMPACT & CONCLUSION
    # =========================================================================
    s6 = prs.slides.add_slide(blank_layout)
    apply_slide_canvas(s6)
    add_header(s6, "Impact & Conclusion", "From Dashboard to Digital HR Teammate", "6", "Transforming passive HR metrics into proactive, autonomous organizational resilience.")

    # Before vs After Comparison (Side by side)
    comp_w = Inches(5.72)
    comp_h = Inches(2.4)
    # Before Card
    draw_card(s6, Inches(0.8), Inches(1.5), comp_w, comp_h, bg_color=C_CARD_BG)
    b_bar = s6.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.5), comp_w, Inches(0.1))
    b_bar.fill.solid()
    b_bar.fill.fore_color.rgb = RGBColor(239, 68, 68)
    b_bar.line.fill.background()

    btb = s6.shapes.add_textbox(Inches(1.0), Inches(1.68), comp_w - Inches(0.4), comp_h - Inches(0.3))
    btf = btb.text_frame
    btf.word_wrap = True
    bp0 = btf.paragraphs[0]
    bp0.text = "BEFORE: Traditional Dashboard Fatigue"
    bp0.font.size = Pt(11)
    bp0.font.bold = True
    bp0.font.color.rgb = RGBColor(239, 68, 68)

    before_flow = [
        "Employee Feedback",
        "Static Dashboard",
        "Manual Interpretation",
        "Slow Decisions",
        "Delayed Follow-up"
    ]
    bp1 = btf.add_paragraph()
    bp1.text = "  ➔  ".join(before_flow)
    bp1.font.size = Pt(9.5)
    bp1.font.color.rgb = C_BODY
    bp1.space_before = Pt(6)

    bp2 = btf.add_paragraph()
    bp2.text = "Result: Burnout goes unaddressed until resignation notices arrive."
    bp2.font.size = Pt(9)
    bp2.font.color.rgb = C_MUTED
    bp2.space_before = Pt(6)

    # After Card
    draw_card(s6, Inches(6.8), Inches(1.5), comp_w, comp_h, bg_color=C_CARD_BG)
    a_bar = s6.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(6.8), Inches(1.5), comp_w, Inches(0.1))
    a_bar.fill.solid()
    a_bar.fill.fore_color.rgb = C_GREEN
    a_bar.line.fill.background()

    atb = s6.shapes.add_textbox(Inches(7.0), Inches(1.68), comp_w - Inches(0.4), comp_h - Inches(0.3))
    atf = atb.text_frame
    atf.word_wrap = True
    ap0 = atf.paragraphs[0]
    ap0.text = "AFTER: PeoplePulse Autonomous Teammate"
    ap0.font.size = Pt(11)
    ap0.font.bold = True
    ap0.font.color.rgb = C_GREEN

    after_flow = [
        "Daily Check-ins",
        "PulseAgent",
        "Autonomous Decisions",
        "Tool Execution",
        "Adaptation",
        "HR Outcome"
    ]
    ap1 = atf.add_paragraph()
    ap1.text = "  ➔  ".join(after_flow)
    ap1.font.size = Pt(9.5)
    ap1.font.color.rgb = C_BODY
    ap1.space_before = Pt(6)

    ap2 = atf.add_paragraph()
    ap2.text = "Result: Immediate diagnosis and automated intervention within minutes."
    ap2.font.size = Pt(9)
    ap2.font.color.rgb = C_GREEN
    ap2.space_before = Pt(6)

    # Three Impact Cards
    impact_items = [
        ("⚡ Faster Identification", "Detects burnout patterns days ahead through live check-in telemetry and NLP.", C_PRIMARY, C_PRIMARY_TINT),
        ("🎯 Targeted Interventions", "Equips managers with personalized 1:1 coaching guides without employee privacy risk.", C_AI_PURPLE, C_PURPLE_TINT),
        ("🔄 Continuous Improvement", "Deploys targeted adaptive surveys to measure recovery and close feedback loops.", C_GREEN, C_GREEN_TINT)
    ]
    imp_w = Inches(3.72)
    imp_h = Inches(1.3)
    for ii, (i_title, i_desc, i_col, i_bg) in enumerate(impact_items):
        ix = Inches(0.8) + ii * Inches(4.0)
        draw_card(s6, ix, Inches(4.1), imp_w, imp_h, bg_color=C_CARD_BG)

        # Top strip
        istrip = s6.shapes.add_shape(MSO_SHAPE.RECTANGLE, ix, Inches(4.1), imp_w, Inches(0.08))
        istrip.fill.solid()
        istrip.fill.fore_color.rgb = i_col
        istrip.line.fill.background()

        itb = s6.shapes.add_textbox(ix + Inches(0.18), Inches(4.25), imp_w - Inches(0.36), imp_h - Inches(0.2))
        itf = itb.text_frame
        itf.word_wrap = True
        ip1 = itf.paragraphs[0]
        ip1.text = i_title
        ip1.font.size = Pt(11)
        ip1.font.bold = True
        ip1.font.color.rgb = i_col
        ip2 = itf.add_paragraph()
        ip2.text = i_desc
        ip2.font.size = Pt(9)
        ip2.font.color.rgb = C_BODY
        ip2.space_before = Pt(3)

    # Final Powerful Statement in Large Text
    fin_card = draw_card(s6, Inches(0.8), Inches(5.6), Inches(11.73), Inches(0.85), bg_color=C_PRIMARY_TINT, border_color=RGBColor(191, 219, 254), border_width=1.5)
    ftb = s6.shapes.add_textbox(Inches(0.9), Inches(5.65), Inches(11.53), Inches(0.75))
    ftf = ftb.text_frame
    ftf.vertical_anchor = MSO_ANCHOR.MIDDLE
    fp = ftf.paragraphs[0]
    fp.text = "“PeoplePulse doesn't just analyze employee wellbeing. PulseAgent helps decide what to do next.”"
    fp.font.size = Pt(13)
    fp.font.bold = True
    fp.font.color.rgb = C_PRIMARY
    fp.alignment = PP_ALIGN.CENTER

    # Footer
    foot_tb = s6.shapes.add_textbox(Inches(0.8), Inches(6.75), Inches(11.73), Inches(0.4))
    foot_tf = foot_tb.text_frame
    foot_p = foot_tf.paragraphs[0]
    foot_p.text = "Tech Zephyr 4.0 | IIT Bhubaneswar    •    PeoplePulse × PulseAgent"
    foot_p.font.size = Pt(10)
    foot_p.font.bold = True
    foot_p.font.color.rgb = C_MUTED
    foot_p.alignment = PP_ALIGN.CENTER

    prs.save(output_path)
    print(f"[SUCCESS] Presentation generated cleanly at: {output_path}")

if __name__ == "__main__":
    build_presentation()
